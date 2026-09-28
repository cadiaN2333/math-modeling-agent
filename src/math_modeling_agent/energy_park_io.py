"""读取并校验能源园区问题的 Excel 曲线数据。"""

from collections.abc import Sequence
from pathlib import Path
import re
from typing import Any

from math import isfinite

from pydantic import ValidationError

from .energy_park_models import (
    EnergyParkCostParameters,
    EnergyParkDataset,
    HourlyProfile,
)


_ATTACHMENT_PREFIXES = {
    1: "附件1：",
    2: "附件2：",
    3: "附件3：",
    4: "附件4：",
    5: "附件5：",
    6: "附件6：",
    7: "附件7：",
    8: "附件8：",
}


def _find_attachment(directory: Path, number: int) -> Path:
    """按附件编号定位唯一的 XLSX 文件。"""

    prefix = _ATTACHMENT_PREFIXES[number]
    candidates = sorted(directory.glob(f"{prefix}*.xlsx"))
    if len(candidates) != 1:
        raise ValueError(
            f"附件目录中应有且仅有一个以 {prefix} 开头的 XLSX 文件，"
            f"实际找到 {len(candidates)} 个"
        )
    return candidates[0]


def _read_workbook_rows(path: Path) -> list[tuple[Any, ...]]:
    """只读 XLSX 的第一个工作表，不执行或保存工作簿内容。"""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            '缺少 Excel 读取依赖；请运行 python -m pip install -e ".[dev,energy]"'
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise ValueError(f"{path.name} 不包含工作表")
        return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def _numeric_cell(
    rows: Sequence[Sequence[Any]],
    *,
    row: int,
    column: int,
    source_name: str,
    expected_row_label: str | None = None,
) -> float:
    """读取一格有限数值，并验证所在行标签。"""

    row_index = row - 1
    column_index = column - 1
    if row_index >= len(rows) or column_index >= len(rows[row_index]):
        raise ValueError(f"{source_name} 缺少单元格第 {row} 行第 {column} 列")
    values = rows[row_index]
    if expected_row_label is not None:
        label = str(values[0] or "")
        if expected_row_label not in label:
            raise ValueError(
                f"{source_name} 第 {row} 行标签应包含“{expected_row_label}”，实际为“{label}”"
            )

    raw_value = values[column_index]
    if isinstance(raw_value, bool) or raw_value is None:
        raise ValueError(f"{source_name} 第 {row} 行第 {column} 列不是数值")
    if isinstance(raw_value, (int, float)):
        result = float(raw_value)
    else:
        match = re.search(r"[-+]?\d+(?:\.\d+)?", str(raw_value))
        if match is None:
            raise ValueError(f"{source_name} 第 {row} 行第 {column} 列不含数值")
        result = float(match.group())
    if not isfinite(result):
        raise ValueError(f"{source_name} 第 {row} 行第 {column} 列必须是有限数值")
    return result


def _text_cell(
    rows: Sequence[Sequence[Any]],
    *,
    row: int,
    column: int,
    source_name: str,
) -> str:
    """读取一格非空文本。"""

    row_index = row - 1
    column_index = column - 1
    if row_index >= len(rows) or column_index >= len(rows[row_index]):
        raise ValueError(f"{source_name} 缺少单元格第 {row} 行第 {column} 列")
    value = rows[row_index][column_index]
    if value is None or not str(value).strip():
        raise ValueError(f"{source_name} 第 {row} 行第 {column} 列不能为空")
    return str(value).strip()


def _parse_cost_parameters(
    attachment_5: Sequence[Sequence[Any]],
    attachment_6: Sequence[Sequence[Any]],
    attachment_7: Sequence[Sequence[Any]],
    attachment_8: Sequence[Sequence[Any]],
) -> EnergyParkCostParameters:
    """按附件行列位置读取成本和电价，并核对行标签。"""

    for row, label in (
        (2, "度电成本"),
        (3, "运维系数"),
        (4, "使用寿命"),
        (5, "效率"),
        (6, "制氢耗电量"),
    ):
        _text_cell(attachment_5, row=row, column=1, source_name="附件5")
        if label not in str(attachment_5[row - 1][0]):
            raise ValueError(f"附件5 第 {row} 行标签应包含“{label}”")
    for row, label in (
        (2, "投资成本"),
        (3, "运维系数"),
        (4, "使用寿命"),
        (5, "效率"),
        (6, "自损耗率"),
        (7, "用电/用氢需求"),
    ):
        if row - 1 >= len(attachment_6) or label not in str(attachment_6[row - 1][0]):
            raise ValueError(f"附件6 第 {row} 行标签应包含“{label}”")

    storage_efficiency_text = _text_cell(
        attachment_6,
        row=5,
        column=2,
        source_name="附件6",
    )
    storage_efficiencies = re.findall(r"\d+(?:\.\d+)?", storage_efficiency_text)
    if len(storage_efficiencies) != 2:
        raise ValueError("附件6 储能效率应分别提供充电和放电效率")

    ammonia_requirements = _text_cell(
        attachment_6,
        row=7,
        column=3,
        source_name="附件6",
    )
    ammonia_energy_match = re.search(
        r"(\d+(?:\.\d+)?)\s*kWh\s*/\s*kgNH3",
        ammonia_requirements,
        re.IGNORECASE,
    )
    ammonia_hydrogen_match = re.search(
        r"(\d+(?:\.\d+)?)\s*kgH2\s*/\s*kgNH3",
        ammonia_requirements,
        re.IGNORECASE,
    )
    if ammonia_energy_match is None or ammonia_hydrogen_match is None:
        raise ValueError("附件6 合成氨装置应提供单位氨耗电和耗氢参数")

    return EnergyParkCostParameters(
        wind_lcoe_yuan_per_kwh=_numeric_cell(
            attachment_5,
            row=2,
            column=2,
            source_name="附件5",
            expected_row_label="度电成本",
        ),
        pv_lcoe_yuan_per_kwh=_numeric_cell(
            attachment_5,
            row=2,
            column=3,
            source_name="附件5",
        ),
        alkaline_om_yuan_per_kwh=_numeric_cell(
            attachment_5,
            row=3,
            column=4,
            source_name="附件5",
            expected_row_label="运维系数",
        ),
        pem_om_yuan_per_kwh=_numeric_cell(
            attachment_5,
            row=3,
            column=5,
            source_name="附件5",
        ),
        alkaline_lifetime_years=int(_numeric_cell(
            attachment_5,
            row=4,
            column=4,
            source_name="附件5",
            expected_row_label="使用寿命",
        )),
        pem_lifetime_years=int(_numeric_cell(
            attachment_5,
            row=4,
            column=5,
            source_name="附件5",
        )),
        alkaline_efficiency_percent=_numeric_cell(
            attachment_5,
            row=5,
            column=4,
            source_name="附件5",
            expected_row_label="效率",
        ),
        pem_efficiency_percent=_numeric_cell(
            attachment_5,
            row=5,
            column=5,
            source_name="附件5",
        ),
        hydrogen_energy_kwh_per_kg=_numeric_cell(
            attachment_5,
            row=6,
            column=4,
            source_name="附件5",
            expected_row_label="制氢耗电量",
        ),
        storage_capex_yuan_per_kwh=_numeric_cell(
            attachment_6,
            row=2,
            column=2,
            source_name="附件6",
            expected_row_label="投资成本",
        ),
        storage_om_yuan_per_kwh=_numeric_cell(
            attachment_6,
            row=3,
            column=2,
            source_name="附件6",
            expected_row_label="运维系数",
        ),
        storage_lifetime_years=int(_numeric_cell(
            attachment_6,
            row=4,
            column=2,
            source_name="附件6",
            expected_row_label="使用寿命",
        )),
        storage_charge_efficiency_percent=float(storage_efficiencies[0]),
        storage_discharge_efficiency_percent=float(storage_efficiencies[1]),
        storage_self_loss_percent=_numeric_cell(
            attachment_6,
            row=6,
            column=2,
            source_name="附件6",
            expected_row_label="自损耗率",
        ),
        ammonia_capex_yuan_per_kg_h2=_numeric_cell(
            attachment_6,
            row=2,
            column=3,
            source_name="附件6",
        ),
        ammonia_om_yuan_per_kwh=_numeric_cell(
            attachment_6,
            row=3,
            column=3,
            source_name="附件6",
        ),
        ammonia_lifetime_years=int(_numeric_cell(
            attachment_6,
            row=4,
            column=3,
            source_name="附件6",
        )),
        ammonia_energy_kwh_per_kg=float(ammonia_energy_match.group(1)),
        ammonia_hydrogen_kg_per_kg=float(ammonia_hydrogen_match.group(1)),
        purchase_price_peak_yuan_per_kwh=_numeric_cell(
            attachment_7,
            row=2,
            column=2,
            source_name="附件7",
            expected_row_label="高峰",
        ),
        purchase_price_flat_yuan_per_kwh=_numeric_cell(
            attachment_7,
            row=3,
            column=2,
            source_name="附件7",
            expected_row_label="平时",
        ),
        purchase_price_valley_yuan_per_kwh=_numeric_cell(
            attachment_7,
            row=4,
            column=2,
            source_name="附件7",
            expected_row_label="低谷",
        ),
        wind_export_price_yuan_per_kwh=_numeric_cell(
            attachment_8,
            row=2,
            column=2,
            source_name="附件8",
            expected_row_label="风电",
        ),
        pv_export_price_yuan_per_kwh=_numeric_cell(
            attachment_8,
            row=3,
            column=2,
            source_name="附件8",
            expected_row_label="光伏",
        ),
    )


def load_energy_park_directory(directory: str | Path) -> EnergyParkDataset:
    """读取题目附件 1—8 并构造经校验的能源园区数据集。"""

    root = Path(directory)
    if not root.is_dir():
        raise ValueError(f"附件目录不存在：{root}")

    paths = {number: _find_attachment(root, number) for number in _ATTACHMENT_PREFIXES}
    tables = {number: _read_workbook_rows(path) for number, path in paths.items()}

    typical_wind = parse_hourly_table(
        tables[2],
        source_name="附件2",
        value_column=1,
        expected_value_header="风电标幺值",
    )
    typical_pv = parse_hourly_table(
        tables[2],
        source_name="附件2",
        value_column=2,
        expected_value_header="光伏标幺值",
    )
    wind_scenarios = [
        parse_hourly_table(
            tables[3],
            source_name="附件3",
            value_column=column,
            expected_value_header=f"风电场景{column}",
        )
        for column in range(1, 7)
    ]
    pv_scenarios = [
        parse_hourly_table(
            tables[4],
            source_name="附件4",
            value_column=column,
            expected_value_header=f"光伏场景{column}",
        )
        for column in range(1, 5)
    ]

    costs = _parse_cost_parameters(
        tables[5],
        tables[6],
        tables[7],
        tables[8],
    )
    try:
        return EnergyParkDataset(
            ordinary_load=parse_hourly_table(
                tables[1],
                source_name="附件1",
                expected_value_header="典型日常规电负荷标幺功率",
            ),
            typical_wind=typical_wind,
            typical_pv=typical_pv,
            wind_scenarios=wind_scenarios,
            pv_scenarios=pv_scenarios,
            costs=costs,
            source_files=[path.name for path in paths.values()],
        )
    except ValidationError as exc:
        raise ValueError(f"能源园区附件数据校验失败：{exc}") from exc


def parse_hourly_table(
    rows: Sequence[Sequence[Any]],
    *,
    source_name: str,
    value_column: int = 1,
    expected_value_header: str | None = None,
) -> HourlyProfile:
    """从含表头的矩阵读取一条 24 小时标幺功率曲线。"""

    if not rows:
        raise ValueError(f"{source_name} 工作表为空")
    if value_column <= 0:
        raise ValueError("功率数据列必须位于时段列之后")
    if len(rows) != 25:
        raise ValueError(f"{source_name} 应有 1 行表头和 24 行小时数据")
    if len(rows[0]) <= value_column:
        raise ValueError(f"{source_name} 表头缺少第 {value_column + 1} 列")
    if expected_value_header is not None:
        actual_header = str(rows[0][value_column] or "")
        if expected_value_header not in actual_header:
            raise ValueError(
                f"{source_name} 第 1 行第 {value_column + 1} 列标题应包含"
                f"“{expected_value_header}”，实际为“{actual_header}”"
            )

    periods: list[str] = []
    values: list[float] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) <= value_column:
            raise ValueError(
                f"{source_name} 第 {row_number} 行缺少第 {value_column + 1} 列数据"
            )
        period, raw_value = row[0], row[value_column]
        if period is None or raw_value is None:
            raise ValueError(f"{source_name} 第 {row_number} 行存在空时段或空数值")
        try:
            numeric_value = float(raw_value)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"{source_name} 第 {row_number} 行的功率值不是数值"
            ) from exc
        periods.append(str(period).strip())
        values.append(numeric_value)

    try:
        return HourlyProfile(periods=periods, values=values)
    except ValidationError as exc:
        raise ValueError(f"{source_name} 曲线数据校验失败：{exc}") from exc
